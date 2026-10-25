import base64
import io
import os
import re
import time
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel
import requests
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

app = FastAPI(
    title="Misykah AI Transcription API",
    description="Layanan AI Filologi Islam untuk Transkripsi & Terjemahan Kitab Arab",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DEFAULT_GEMINI_KEY = os.getenv("GEMINI_API_KEY", "AQ.Ab8RN6JNP6bqsdD3NReSWf7ph26YBuwFpmLGmq3wjYhxK8YRLw")
GEMINI_MODELS = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-1.5-flash"]

# ── Helper Functions ──────────────────────────────────────────

def is_heading_line(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 120:
        return False
    if s.startswith("#") or s.startswith("**[") or s.startswith("---"):
        return True
    keywords = ["البرهان", "فصل", "باب", "المبحث", "المطلب", "مسألة", "المقدمة", "الخاتمة", "فهرس", "قاعدة", "تنبيه"]
    for kw in keywords:
        if s.startswith(kw) or re.match(r"^(?:فصل|باب|المبحث|المطلب|مسألة|البرهان|قاعدة)\s+[\w\d\u0660-\u0669]", s):
            return True
    return False

def is_footnote_line(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    return bool(re.match(r"^[\(\[]?[\d\u0660-\u0669]+[\)\]\.\:\-]", s))

def set_paragraph_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def parse_gemini_kitab_output(text: str):
    arabic_lines = []
    translation_lines = []
    footnotes = []
    headings = []

    is_translating = False
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue

        if is_footnote_line(line):
            is_translating = False
            footnotes.append(line)
            continue

        if line.startswith("[TERJEMAHAN") or line.startswith("[Terjemahan"):
            is_translating = True
            continue
        elif is_translating and any('\u0600' <= c <= '\u06FF' for c in line) and not re.search(r"[A-Za-z]", line):
            is_translating = False

        if is_translating:
            clean_tr = re.sub(r"^Arti\s*:\s*", "", line)
            translation_lines.append(clean_tr)
        else:
            if is_heading_line(line):
                headings.append(re.sub(r"[*#_]", "", line).strip())
            arabic_lines.append(line)

    return {
        "arabic_text": "\n".join(arabic_lines),
        "translation": "\n".join(translation_lines),
        "footnotes": footnotes,
        "headings": headings,
        "raw_text": text
    }

def call_gemini_vision(img_bytes: bytes, include_translation: bool = True, api_key: str = DEFAULT_GEMINI_KEY) -> str:
    b64_img = base64.b64encode(img_bytes).decode("utf-8")

    if include_translation:
        prompt = """
        Ini adalah halaman buku/kitab berbahasa Arab.
        Tolong transkripsikan teks Arab dan berikan terjemahan bahasa Indonesia yang ilmiah:
        1. Salin seluruh teks Arab persis sesuai aslinya dalam format teks Arab Unicode (UTF-8) lengkap dengan harakat/tasykil jika ada.
        2. Tepat di bawah setiap ayat, hadits, atau paragraf teks Arab, sertakan terjemahan bahasa Indonesia yang akurat, ilmiah, dan luwes dengan format:
           [TERJEMAHAN]
           (terjemahan bahasa Indonesia di sini)
        3. Pisahkan antara judul bab, hadits, ayat Al-Qur'an, dan catatan kaki secara rapi.
        4. Dilarang memberikan komentar pengantar tambahan, keluarkan murni teks Arab dan terjemahannya saja.
        """
    else:
        prompt = """
        Ini adalah halaman buku/kitab berbahasa Arab.
        Tolong transkripsikan seluruh teks bahasa Arab pada gambar halaman ini dengan presisi tinggi dan sangat teliti:
        1. Salin seluruh teks persis sesuai aslinya dalam format teks Arab Unicode (UTF-8).
        2. Pertahankan harakat (tasykil) jika ada di naskah asli.
        3. Pisahkan antara judul, teks utama, hadits, dan catatan kaki secara rapi.
        4. Dilarang menerjemahkan atau memberikan komentar/mukadimah, keluarkan murni teks Arab hasil transkripsi saja.
        """

    payload = {
        "contents": [{
            "parts": [
                {"inlineData": {"mimeType": "image/jpeg", "data": b64_img}},
                {"text": prompt}
            ]
        }],
        "generationConfig": {
            "temperature": 0.1,
            "maxOutputTokens": 4096
        }
    }
    headers = {"Content-Type": "application/json"}

    last_error = "Gagal memproses setelah beberapa kali percobaan ke Gemini."
    for model_name in GEMINI_MODELS:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
        for attempt in range(2):
            try:
                resp = requests.post(url, headers=headers, json=payload, timeout=50)
                if resp.status_code == 200:
                    res_data = resp.json()
                    candidates = res_data.get("candidates", [])
                    if candidates:
                        parts = candidates[0].get("content", {}).get("parts", [])
                        if parts:
                            return parts[0].get("text", "").strip()
                elif resp.status_code in (429, 503):
                    time.sleep(2 * (attempt + 1))
                    continue
                else:
                    last_error = f"Gemini Error ({resp.status_code}): {resp.text[:180]}"
                    break
            except Exception as e:
                last_error = str(e)
                time.sleep(1)

    raise HTTPException(status_code=500, detail=last_error)

# ── API Endpoints ─────────────────────────────────────────────

@app.get("/api/health")
@app.get("/health")
def health_check():
    return {
        "status": "online",
        "service": "(مشكاة) Misykah Cloud Transcription AI",
        "engine": "Google Gemini Vision"
    }

@app.post("/api/transcribe-canvas")
@app.post("/transcribe-canvas")
async def transcribe_canvas(
    file: UploadFile = File(..., description="Tangkapan layar halaman buku"),
    include_translation: bool = Form(True),
    strip_tashkeel: bool = Form(False)
):
    img_bytes = await file.read()
    raw_res = call_gemini_vision(img_bytes, include_translation=include_translation)

    if strip_tashkeel:
        raw_res = re.sub(r"[\u064B-\u065F\u0670]", "", raw_res)

    parsed = parse_gemini_kitab_output(raw_res)
    return {
        "success": True,
        "data": parsed
    }

class ExportWordRequest(BaseModel):
    title: str = "Transkripsi Kitab"
    author: Optional[str] = "Santri Misykah"
    pages: list = []

@app.post("/api/export-word")
@app.post("/export-word")
async def export_word_doc(req: ExportWordRequest):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title
    p_title = doc.add_paragraph()
    set_paragraph_rtl(p_title)
    r_title = p_title.add_run(req.title)
    r_title.font.name = "Traditional Arabic"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(14, 61, 38)

    # Pages
    for p_item in req.pages:
        p_num = p_item.get("page", 1)
        arabic = p_item.get("arabic_text", "")
        translation = p_item.get("translation", "")
        footnotes = p_item.get("footnotes", [])

        # Heading Page
        p_h = doc.add_paragraph()
        set_paragraph_rtl(p_h)
        r_h = p_h.add_run(f"--- [ صفحة {p_num} ] ---")
        r_h.font.name = "Traditional Arabic"
        r_h.font.size = Pt(13)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(180, 140, 20)

        # Arabic Text
        if arabic:
            for line in arabic.splitlines():
                if not line.strip():
                    continue
                p_ar = doc.add_paragraph()
                set_paragraph_rtl(p_ar)
                r_ar = p_ar.add_run(line)
                r_ar.font.name = "Traditional Arabic"
                r_ar.font.size = Pt(16)
                r_ar.font.color.rgb = RGBColor(10, 10, 10)

        # Translation
        if translation:
            p_tr_label = doc.add_paragraph()
            r_tr_l = p_tr_label.add_run("[ Terjemahan Indonesia ]")
            r_tr_l.font.name = "Calibri"
            r_tr_l.font.size = Pt(10.5)
            r_tr_l.font.bold = True
            r_tr_l.font.color.rgb = RGBColor(14, 61, 38)

            for line in translation.splitlines():
                if not line.strip():
                    continue
                p_tr = doc.add_paragraph()
                r_tr = p_tr.add_run(line)
                r_tr.font.name = "Calibri"
                r_tr.font.size = Pt(11)

        # Footnotes
        if footnotes:
            doc.add_paragraph("________________________________________")
            for fn in footnotes:
                p_fn = doc.add_paragraph()
                set_paragraph_rtl(p_fn)
                r_fn = p_fn.add_run(fn)
                r_fn.font.name = "Traditional Arabic"
                r_fn.font.size = Pt(11)
                r_fn.font.color.rgb = RGBColor(90, 90, 90)

    # Save to memory stream
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    return Response(
        content=bio.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="Transkripsi_{req.title}.docx"'}
    )
